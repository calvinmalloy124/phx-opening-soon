from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF", size=10)          # inputs
BLACK = Font(name=FONT, color="000000", size=10)         # formulas
GREEN = Font(name=FONT, color="008000", size=10)         # cross-sheet links
BOLD = Font(name=FONT, bold=True, size=10)
TITLE = Font(name=FONT, bold=True, size=14, color="1F3864")
SUB = Font(name=FONT, italic=True, size=9, color="666666")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TOTAL_FILL = PatternFill("solid", fgColor="E7E6E6")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

CUR = '$#,##0.00;($#,##0.00);-'
CUR0 = '$#,##0;($#,##0);-'
PCT = '0.0%;(0.0%);-'
NUM = '#,##0.00;(#,##0.00);-'
INT = '#,##0;(#,##0);-'

DISCLAIMER = ("For planning and educational use only. Not financial, legal, or tax advice. "
              "Every number is an estimate you should verify for your own business.")


def new_wb():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def title(ws, text, sub=None, row=1):
    ws.cell(row=row, column=1, value=text).font = TITLE
    if sub:
        ws.cell(row=row + 1, column=1, value=sub).font = SUB
    return row + (3 if sub else 2)


def legend(ws, row, col=1):
    c = ws.cell(row=row, column=col, value="How to use: type your numbers in the yellow cells (blue text). Everything else calculates.")
    c.font = SUB
    return row + 1


def header(ws, row, headers, col=1, widths=None):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(col + i)].width = w
    ws.row_dimensions[row].height = 30
    return row + 1


def inp(ws, row, col, value, fmt=None, key=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = BLUE
    c.fill = KEY_FILL if key else INPUT_FILL
    c.border = BOX
    if fmt:
        c.number_format = fmt
    return c


def fx(ws, row, col, formula, fmt=None, bold=False, total=False, link=False):
    c = ws.cell(row=row, column=col, value=formula)
    c.font = Font(name=FONT, bold=bold, size=10, color="008000" if link else "000000")
    c.border = BOX
    if total:
        c.fill = TOTAL_FILL
    if fmt:
        c.number_format = fmt
    return c


def label(ws, row, col, text, bold=False, wrap=False):
    c = ws.cell(row=row, column=col, value=text)
    c.font = BOLD if bold else BLACK
    c.border = BOX
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def note(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.font = SUB
    c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def readme(wb, product, steps, notes=()):
    ws = wb.create_sheet("Start Here", 0)
    r = title(ws, product, "Restaurant & Bar Operator Toolkit")
    ws.cell(row=r, column=1, value="Setup").font = BOLD
    r += 1
    for i, s in enumerate(steps, 1):
        note(ws, r, 1, f"{i}. {s}")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Conventions").font = BOLD
    r += 1
    for t in ["Yellow cells with blue text are inputs. Type over the sample numbers.",
              "Black text is a formula. Leave it alone unless you know what you're changing.",
              "Green text pulls from another tab.",
              "Works in Excel 2016+ and Google Sheets (File > Import > Upload > Replace). No macros."] + list(notes):
        note(ws, r, 1, "• " + t)
        r += 1
    r += 1
    note(ws, r, 1, DISCLAIMER)
    ws.column_dimensions["A"].width = 110
    return ws


def freeze(ws, cell):
    ws.freeze_panes = cell
